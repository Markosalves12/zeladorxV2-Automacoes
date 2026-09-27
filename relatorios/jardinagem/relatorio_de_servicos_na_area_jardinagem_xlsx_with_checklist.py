import openpyxl
from servicos.headers_report_jardinagem import headers_report_services, headers_report_schedules
from servicos.utils_jardinagem import colect_dados_fato_servico_jardinagem, colect_dados_agendamentos_jardinagem
from checklists.models import CheckListJardinagem  # Import do modelo de checklist
from django.http import HttpResponse
from utils.utils import generate_id_random, filtrar_unicos
from datetime import datetime
from permissionscontrol.utils import verify_login
from django.shortcuts import redirect
from relatorios.utils_xlsx import adicionar_cabecalhos, adicionar_dados

def exportar_relatorio_de_serivos_na_area_Jardinagem_excel_with_checklist(request, userid, id_random, DataDeInicio, DataDeConclusao, Areas,
                                                                         TipoServico, ServicosEscalados, ColaboradoresEscalados, type):
    block = verify_login(request=request, userid=userid)

    if block:
        return redirect('logout')

    # Configuração inicial do workbook
    wb = openpyxl.Workbook()
    ws_acompanhamento = wb.active
    ws_acompanhamento.title = "Relatório de acompanhamento"
    ws_agendadamentos = wb.create_sheet(title="Relatório de agendamentos jardinagem") # Cria a segunda aba
    ws_checklists = wb.create_sheet(title="Dados do Checklist")  # Cria a terceira aba

    # Converte os parâmetros de entrada
    DataDeInicio = datetime.strptime(DataDeInicio, '%Y-%m-%dT%H:%M') if DataDeInicio and DataDeInicio != "None" else 'None'
    DataDeConclusao = datetime.strptime(DataDeConclusao, '%Y-%m-%dT%H:%M') if DataDeConclusao and DataDeConclusao != "None" else 'None'
    ServicosEscalados = ServicosEscalados.split(',')
    ColaboradoresEscalados = ColaboradoresEscalados.split(',')


    # Coleta os dados
    dados = colect_dados_fato_servico_jardinagem(
        request=request,
        userid=userid,
        DataDeInicio=DataDeInicio,
        DataDeConclusao=DataDeConclusao,
        TipoServico=TipoServico,
        Areas=Areas,
        ServicosEscalados=ServicosEscalados,
        ColaboradoresEscalados=ColaboradoresEscalados,
        status=['Concluido'],
    )

    # Filtra os dados conforme o tipo
    if type == 'catalogo_de_servicos':
        dados = filtrar_unicos(dados, Servico__ServicosEscalados__id_random=id_random)
    elif type == 'configuracao':
        dados = dados.filter(Servico__id_configuracao=id_random)

    elif type == 'areas':
        dados = dados.filter(Servico__Areas__id_random=id_random)

    elif type == 'gerente':
        dados = dados.filter(Gerente__id_random=id_random)

    # Definindo as colunas com seus tipos, conforme o código original
    colunas = [
        ('id_acompanhamento', None),
        ('id_random_acompanhamento', None),
        ('id_agendamento', None),
        ('id_random_agendamento', None),
        ('colaboradores_chamados_id', None),
        ('colaboradores_chamados_id_random', None),
        ('colaboradores_chamados', None),
        ('data_hora_chegada', 'data'),  # Tratamento especial para data
        ('data_hora_retorno', 'data'),  # Tratamento especial para data
        ('tempo_na_area', None),
    ]

    # Chamando as funções
    adicionar_cabecalhos(ws_acompanhamento, headers_report_services)
    adicionar_dados(ws_acompanhamento, dados, colunas, headers_report_services)

    dados_de_agendamentos = colect_dados_agendamentos_jardinagem(
        request=request,
        userid=userid,
        DataDeInicio=DataDeInicio,
        DataDeConclusao=DataDeConclusao,
        ServicosEscalados=ServicosEscalados,
        ColaboradoresEscalados=ColaboradoresEscalados,
        TipoServico=TipoServico,
        Areas=Areas,
        status=['Concluido'],
    )

    # Obter todos os IDs únicos de uma vez
    ids_random_agendamentos = {dado.id_random_agendamento for dado in dados}

    # Filtrar todos os dados de uma vez
    dados_agendamentos = dados_de_agendamentos.filter(id_random_agendamento__in=ids_random_agendamentos)

    # Definindo as colunas com seus tipos, conforme o código original
    colunas = [
        ('tipodeempresa', None),
        ('empresaprestadora', None),
        ('id_agendamento', None),
        ('id_random_agendamento', None),
        ('id_config', None),
        ('tipo_agendamento', None),
        ('descricao_do_servico', None),
        ('colaboradores_chamados', None),
        ('colaboradores_chamados_id', None),
        ('colaboradores_chamados_id_random', None),
        ('servicos_solicitados', None),
        ('servicos_solicitados_id', None),
        ('servicos_solicitados_id_random', None),
        ('data_de_inicio', 'data'),  # Tratamento especial para data
        ('data_de_conclusao', 'data'),  # Tratamento especial para data
        ('antes', None),
        ('depois', None),
        ('status_servico', None),
        ('area_atendida', None),
        ('area_atendida_id', None),
        ('id_random_area', None),
        ('periodicidade_de_retorno', None),
        ('area_total', None),
        ('tipo_vegetacao', None),
        ('tipo_terreno', None),
        ('localidade', None),
        ('lat', None),
        ('long', None),
        ('unidade', None)
    ]

    # Chamando as funções
    adicionar_cabecalhos(ws_agendadamentos, headers_report_schedules)
    adicionar_dados(ws_agendadamentos, dados_agendamentos, colunas, headers_report_schedules)

    # Define cabeçalhos para a aba de checklists
    checklist_headers = [
        "Agendamento.ID", "Agendamento.Descrição", "Checklist.Descrição",
        "Checklist.Foto", "Checklist.AtualizadoEm", "Checklist.Status"
    ]

    dados_checklist = CheckListJardinagem.objects.filter(servico_agendado__id_random__in=ids_random_agendamentos)

    colunas_checklist = [
        ('servico_agendado.id', 'nested'),  # Atributo aninhado
        ('servico_agendado.DescricaoDoServico', 'nested'),  # Atributo aninhado
        ('descricao', None),
        ('foto_comprovacao', 'url'),  # URL ou "N/A"
        ('atualizado_em', 'datetime_str'),  # Formato de data/hora
        ('status', None)
    ]

    adicionar_cabecalhos(ws_checklists, checklist_headers)
    adicionar_dados(ws_checklists, dados_checklist, colunas_checklist, checklist_headers)

    # Configura a resposta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="relatorio_de_servicos_Concluido_{generate_id_random()}.xlsx"'

    # Salva o Excel na resposta
    wb.save(response)

    return response